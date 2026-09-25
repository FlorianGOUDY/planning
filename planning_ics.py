#!/usr/bin/env python3
"""
Planning ITECH 1 (Excel) + fichier des salles (Excel) -> calendriers .ics

Usage:
    python planning_ics.py --download --tous --outdir docs
        Download the 2 public Google Sheets and generate all calendars.

    python planning_ics.py PLANNING.xlsx SALLES.xlsx --tous
        Generate all calendars from local Excel files.

    python planning_ics.py PLANNING.xlsx SALLES.xlsx --groupe G2
        Generate the complete G2 calendar.

    python planning_ics.py PLANNING.xlsx SALLES.xlsx --groupe G2-A
        Generate the G2-A calendar.

    python planning_ics.py PLANNING.xlsx SALLES.xlsx --groupe G2-B
        Generate the G2-B calendar.

Dependency:
    pip install openpyxl
"""

import argparse
import datetime as dt
import hashlib
import os
import re
import unicodedata
import urllib.request

from collections import defaultdict
from openpyxl import load_workbook


# =====================================================================
# CONFIGURATION
# =====================================================================

FIRST_MONDAY = dt.date(2026, 8, 24)
FIRST_COL = 3                          # Column C
COLS_PER_WEEK = 9                      # 3 groups x 3 columns
FIRST_ROW = 5                          # Monday 08:00
ROWS_PER_DAY = 21                      # 08:00 -> 18:30
DAY_START_MIN = 8 * 60

# Each group occupies 3 columns.
#
# Example:
#
# G1 = columns 0..2
# G2 = columns 3..5
# G3 = columns 6..8
#
GROUP_COLS = {
    "G1": (0, 2),
    "G2": (3, 5),
    "G3": (6, 8),
}

# The first two columns are normally A and B.
# The third column is the common/shared column.
#
# The code also tries to detect the A/B labels directly from the
# spreadsheet header, so the layout is not hard-coded unnecessarily.
DEFAULT_SUBGROUP_COLS = {
    "A": 0,
    "B": 1,
}

ROOM_COLS = range(3, 28)               # C..AA in "Salles de cours"

TZ = "Europe/Paris"

PLANNING_ID = "1ZfgKO-68K1ioRlTjY2x6Z5C8sH1TpJ19nxcdbMMIyPI"
SALLES_ID = "1es1RFXRtJGeyL8MzFaw6fyhEzXDbumx6GTGYiUxBSAg"

MIN_EVENTS = 50


# =====================================================================
# TEXT NORMALIZATION
# =====================================================================

STOP = {
    "td", "tp", "ds", "cm", "de", "des", "du", "la", "le", "les",
    "l", "d", "a", "au", "et", "en",
    "itech", "1", "2", "3", "e", "learning",
    "initiation", "introduction"
}

SYN = {
    "legislation": "droit",
    "travail": "droit",
    "economique": "eco",
    "economie": "eco",
    "resistance": "rdm",
    "materiaux": "rdm",
    "fluides": "flu",
    "mecanique": "meca",
    "thermodynamique": "thermo",
    "organique": "orga",
    "incertitudes": "incertitude",
    "analyses": "analyse",
    "instrumentales": "instrumentale",
    "scientifiques": "bsi",
    "bases": "bsi",
    "ingenieur": "bsi",
    "outils": "oin",
    "numeriques": "oin",
    "informatiques": "oin",
    "generale": "chimie",
    "profil": "profil",
    "parole": "parole",
}


def norm(s):
    """
    Normalize text for comparisons.
    """
    s = (
        unicodedata
        .normalize("NFKD", str(s))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )

    s = re.sub(r"\s*&\s*", "&", s)
    return re.sub(r"[^a-z0-9&]+", " ", s).strip()


def tokens(title):
    """
    Extract normalized semantic tokens from a title.
    """
    out = set()

    for word in norm(title).split():
        if word in STOP or word.isdigit():
            continue

        out.add(SYN.get(word, word))

    return out


def tok_match(a, b):
    """
    Compare token sets while allowing prefix matches.
    """
    return any(
        x == y
        or (
            len(x) >= 3
            and len(y) >= 3
            and (x.startswith(y) or y.startswith(x))
        )
        for x in a
        for y in b
    )


# =====================================================================
# TEACHERS / CONTENT
# =====================================================================

TEACHER_RE = re.compile(
    r"^(?:[A-Z]{1,2}\.\s?|[A-Z]\s)"
    r"([A-ZÀ-Ý][a-zà-ÿ'\-]{2,})"
)


def teachers(lines):
    """
    Detect teacher surnames in lines such as:
        J.Champliaud
        AC.Besson
        L Scalone
    """
    out = set()

    for line in lines:
        match = TEACHER_RE.match(line.strip())

        if match:
            out.add(norm(match.group(1)))

    return out


def content_lines(lines):
    """
    Keep course-related lines:
    - no teacher names
    - no counters such as 1/3
    - no lines containing only numbers/dashes
    """
    return [
        line
        for line in lines
        if not TEACHER_RE.match(line)
        and not re.fullmatch(r"[-\s\d/]+", line)
    ]


def group_marker(text):
    """
    Detect explicit group markers in room information.

    Examples:
        G1
        G2
        Groupe 3
        grp2
    """
    match = re.search(
        r"\b(?:g|grp|groupe)\s*([123])",
        norm(text),
    )

    return f"G{match.group(1)}" if match else None


# =====================================================================
# WEEK DATES
# =====================================================================

def week_starts(ws, n=60):
    """
    Read Monday dates from row 2.

    Invalid dates such as the known 2024 value in C2 are ignored.
    """
    raw = {}

    for k in range(n):
        value = ws.cell(
            2,
            FIRST_COL + COLS_PER_WEEK * k
        ).value

        if (
            isinstance(value, dt.datetime)
            and value.year >= 2025
            and value.weekday() == 0
        ):
            raw[k] = value.date()

    if raw:
        first_index = min(raw)
        anchor = (
            min(raw.items())[1]
            - dt.timedelta(days=7 * first_index)
        )
    else:
        anchor = FIRST_MONDAY

    return [
        raw.get(
            k,
            anchor + dt.timedelta(days=7 * k)
        )
        for k in range(n)
    ]


# =====================================================================
# SUBGROUP DETECTION
# =====================================================================

def detect_subgroup_columns(ws):
    """
    Detect which columns inside each group correspond to A and B.

    The normal layout is:

        Gx
        A   B   [common]

    Therefore:

        first column  -> A
        second column -> B
        third column  -> common/shared

    The function first tries to find explicit "A" and "B" labels
    in the header area. If they cannot be found, it falls back to
    the standard first-column=A / second-column=B layout.
    """

    subgroup_columns = {}

    # ---------------------------------------------------------------
    # Default mapping
    # ---------------------------------------------------------------

    for group, (start, end) in GROUP_COLS.items():
        subgroup_columns[group] = {
            "A": start,
            "B": start + 1,
            "COMMON": set(range(start, end + 1)),
        }

    # ---------------------------------------------------------------
    # Try to detect explicit A/B labels in the header.
    #
    # We intentionally only inspect the rows before FIRST_ROW,
    # because those are the header rows and not actual courses.
    # ---------------------------------------------------------------

    header_max_row = max(1, FIRST_ROW - 1)

    for group, (start, end) in GROUP_COLS.items():

        found_a = None
        found_b = None

        for row in range(1, header_max_row + 1):

            for relative_col in range(start, end + 1):
                absolute_col = FIRST_COL + relative_col

                value = ws.cell(
                    row,
                    absolute_col
                ).value

                if value is None:
                    continue

                value_norm = norm(value)

                if value_norm == "a":
                    found_a = relative_col

                elif value_norm == "b":
                    found_b = relative_col

        # Only override the defaults when both labels were found.
        if found_a is not None and found_b is not None:
            subgroup_columns[group]["A"] = found_a
            subgroup_columns[group]["B"] = found_b

    return subgroup_columns


def subgroups_for_columns(
    group,
    p1,
    p2,
    subgroup_columns,
):
    """
    Determine which subgroups an event belongs to.

    Rules:

    A column only:
        -> A

    B column only:
        -> B

    Common/shared column:
        -> A + B

    A + B:
        -> A + B

    Entire group:
        -> A + B
    """

    if group not in GROUP_COLS:
        return set()

    group_start, group_end = GROUP_COLS[group]

    # Limit the event to this group's columns.
    overlap_start = max(p1, group_start)
    overlap_end = min(p2, group_end)

    if overlap_start > overlap_end:
        return set()

    info = subgroup_columns[group]

    a_col = info["A"]
    b_col = info["B"]

    covered = set(
        range(overlap_start, overlap_end + 1)
    )

    result = set()

    # A is present.
    if a_col in covered:
        result.add("A")

    # B is present.
    if b_col in covered:
        result.add("B")

    # Any additional / common column belongs to both A and B.
    common_columns = covered - {a_col, b_col}

    if common_columns:
        result.update({"A", "B"})

    return result


# =====================================================================
# PLANNING READER
# =====================================================================

def read_planning(path):
    """
    Read the planning Excel file and generate event objects.

    Each event receives:

        groups:
            ["G1"], ["G2"], etc.

        subgroups:
            {"A"}
            {"B"}
            {"A", "B"}

        partial:
            True when the event does not cover the whole group.
    """

    ws = load_workbook(
        path,
        data_only=True
    ).worksheets[0]

    starts = week_starts(ws)

    merged = {
        (m.min_row, m.min_col): m
        for m in ws.merged_cells.ranges
    }

    subgroup_columns = detect_subgroup_columns(ws)

    events = []

    max_row = (
        FIRST_ROW
        + 5 * ROWS_PER_DAY
        - 1
    )

    for row in ws.iter_rows(
        min_row=FIRST_ROW,
        max_row=max_row,
        min_col=FIRST_COL,
    ):

        for cell in row:

            if (
                cell.value is None
                or not str(cell.value).strip()
            ):
                continue

            merge = merged.get(
                (cell.row, cell.column)
            )

            if merge:
                r2 = merge.max_row
                c2 = merge.max_col
            else:
                r2 = cell.row
                c2 = cell.column

            # -------------------------------------------------------
            # Week / column position
            # -------------------------------------------------------

            week = (
                cell.column - FIRST_COL
            ) // COLS_PER_WEEK

            p1 = (
                cell.column - FIRST_COL
            ) % COLS_PER_WEEK

            p2 = (
                c2 - FIRST_COL
            ) % COLS_PER_WEEK

            # A merged cell can theoretically extend across
            # multiple weeks. Keep it inside the current week.
            if c2 - cell.column >= COLS_PER_WEEK:
                p2 = COLS_PER_WEEK - 1

            # -------------------------------------------------------
            # Day / row position
            # -------------------------------------------------------

            day, offset = divmod(
                cell.row - FIRST_ROW,
                ROWS_PER_DAY
            )

            n_rows = (
                r2 - cell.row + 1
            )

            # -------------------------------------------------------
            # Groups
            # -------------------------------------------------------

            groups = [
                group
                for group, (a, b) in GROUP_COLS.items()
                if p1 <= b and p2 >= a
            ]

            full_groups = [
                group
                for group, (a, b) in GROUP_COLS.items()
                if p1 <= a and p2 >= b
            ]

            partial = any(
                group not in full_groups
                for group in groups
            )

            # -------------------------------------------------------
            # Subgroups
            # -------------------------------------------------------

            subgroup_map = {}

            for group in groups:

                subgroup_map[group] = subgroups_for_columns(
                    group,
                    p1,
                    p2,
                    subgroup_columns,
                )

            # -------------------------------------------------------
            # Date
            # -------------------------------------------------------

            base_date = (
                starts[week]
                + dt.timedelta(days=day)
            )

            text = str(cell.value).strip()

            # -------------------------------------------------------
            # Full-day event
            # -------------------------------------------------------

            if (
                offset == 0
                and n_rows >= ROWS_PER_DAY
            ):

                ndays = max(
                    1,
                    min(
                        5 - day,
                        round(
                            n_rows / ROWS_PER_DAY
                        )
                    )
                )

                events.append(
                    dict(
                        date=base_date,
                        end_date=(
                            base_date
                            + dt.timedelta(days=ndays)
                        ),
                        allday=True,
                        text=text,
                        groups=groups,
                        subgroups=subgroup_map,
                        partial=partial,
                    )
                )

                continue

            # -------------------------------------------------------
            # Normal timed event
            # -------------------------------------------------------

            r2 = min(
                r2,
                FIRST_ROW
                + (day + 1) * ROWS_PER_DAY
                - 1
            )

            start = (
                DAY_START_MIN
                + 30 * offset
            )

            end = (
                DAY_START_MIN
                + 30 * (
                    r2
                    - (
                        FIRST_ROW
                        + day * ROWS_PER_DAY
                    )
                    + 1
                )
            )

            # -------------------------------------------------------
            # Explicit time in the cell text.
            #
            # Example:
            #   13h30 - 15h30
            #
            # This takes precedence over the spreadsheet block.
            # -------------------------------------------------------

            time_match = re.search(
                r"(\d{1,2})\s*h\s*(\d{2})?\s*-\s*"
                r"(\d{1,2})\s*h\s*(\d{2})?",
                text,
            )

            if time_match:

                start2 = (
                    int(time_match.group(1)) * 60
                    + int(time_match.group(2) or 0)
                )

                end2 = (
                    int(time_match.group(3)) * 60
                    + int(time_match.group(4) or 0)
                )

                if (
                    6 * 60 <= start2 < end2 <= 22 * 60
                ):
                    start = start2
                    end = end2

            events.append(
                dict(
                    date=base_date,
                    start=start,
                    end=end,
                    allday=False,
                    text=text,
                    groups=groups,
                    subgroups=subgroup_map,
                    partial=partial,
                )
            )

    return events


# =====================================================================
# ROOM READER
# =====================================================================

def read_rooms(path):
    """
    Read the "Salles de cours" sheet.
    """

    ws = load_workbook(
        path,
        data_only=True
    )["Salles de cours"]

    headers = {
        column: str(
            ws.cell(3, column).value or ""
        ).split("\n")[0].strip()
        for column in ROOM_COLS
    }

    merged = {
        (m.min_row, m.min_col): m
        for m in ws.merged_cells.ranges
    }

    output = []

    for row in range(4, ws.max_row + 1):

        date_value = ws.cell(
            row,
            1
        ).value

        if not isinstance(
            date_value,
            dt.datetime
        ):
            continue

        for room_row in range(
            row,
            row + ROWS_PER_DAY
        ):

            for column in ROOM_COLS:

                value = ws.cell(
                    room_row,
                    column
                ).value

                if (
                    value is None
                    or not str(value).strip()
                ):
                    continue

                merge = merged.get(
                    (room_row, column)
                )

                room_row_end = min(
                    merge.max_row
                    if merge
                    else room_row,
                    row + ROWS_PER_DAY - 1,
                )

                output.append(
                    dict(
                        date=date_value.date(),
                        start=(
                            DAY_START_MIN
                            + 30 * (room_row - row)
                        ),
                        end=(
                            DAY_START_MIN
                            + 30 * (
                                room_row_end - row + 1
                            )
                        ),
                        room=headers[column],
                        text=str(value).strip(),
                    )
                )

    return output


# =====================================================================
# PLANNING <-> ROOMS MATCHING
# =====================================================================

def find_rooms(
    event,
    group,
    rooms_by_date,
):
    """
    Match a planning event with rooms.
    """

    if event["allday"]:
        return []

    lines = [
        line.strip()
        for line in event["text"].split("\n")
        if line.strip()
    ]

    if not lines:
        return []

    if norm(lines[0]).startswith("e learning"):
        return []

    content = content_lines(lines)

    title_tokens = (
        set().union(
            *[
                tokens(line)
                for line in content[:2]
            ]
        )
        if content
        else set()
    )

    if not title_tokens:
        return []

    event_teachers = teachers(lines)

    found = []

    for room in rooms_by_date.get(
        event["date"],
        []
    ):

        room_lines = [
            line.strip()
            for line in room["text"].split("\n")
            if line.strip()
        ]

        # -----------------------------------------------------------
        # Keep only ITECH 1 rooms or matching teachers.
        # -----------------------------------------------------------

        if not re.search(
            r"itech\s*1",
            norm(room["text"])
        ) and not (
            event_teachers
            & teachers(room_lines)
        ):
            continue

        # -----------------------------------------------------------
        # Time overlap
        # -----------------------------------------------------------

        overlap = (
            min(event["end"], room["end"])
            - max(event["start"], room["start"])
        )

        if overlap < (
            0.6
            * min(
                event["end"] - event["start"],
                room["end"] - room["start"],
            )
        ):
            continue

        # -----------------------------------------------------------
        # Course title matching
        # -----------------------------------------------------------

        room_content = content_lines(
            room_lines
        )

        room_tokens = (
            set().union(
                *[
                    tokens(line)
                    for line in room_content[:1]
                ]
            )
            if room_content
            else set()
        )

        if not tok_match(
            title_tokens,
            room_tokens
        ):
            continue

        # -----------------------------------------------------------
        # Explicit group marker
        # -----------------------------------------------------------

        marker = group_marker(
            room["text"]
        )

        if (
            marker
            and group
            and marker != group
        ):
            continue

        # -----------------------------------------------------------
        # Teacher matching
        # -----------------------------------------------------------

        room_teachers = teachers(
            room_lines
        )

        if (
            event_teachers
            and room_teachers
            and not (
                event_teachers
                & room_teachers
            )
        ):
            continue

        found.append(
            dict(
                room=room["room"],
                teacher=", ".join(
                    sorted(
                        teacher.title()
                        for teacher in room_teachers
                    )
                )
                if room_teachers
                else "",
                marked=bool(marker),
            )
        )

    # ---------------------------------------------------------------
    # If some rooms explicitly contain the correct group,
    # discard unmarked matches.
    # ---------------------------------------------------------------

    if any(
        item["marked"]
        for item in found
    ):
        found = [
            item
            for item in found
            if item["marked"]
        ]

    # ---------------------------------------------------------------
    # Remove duplicate room names.
    # ---------------------------------------------------------------

    unique = []
    seen = set()

    for item in found:

        if item["room"] in seen:
            continue

        seen.add(item["room"])
        unique.append(item)

    return unique


# =====================================================================
# ICS HELPERS
# =====================================================================

def esc(value):
    """
    Escape an ICS field.
    """
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\n", "\\n")
    )


def fold(line):
    """
    Fold long UTF-8 ICS lines.
    """
    data = line.encode("utf-8")
    output = []

    while len(data) > 74:

        cut = 74

        while (
            cut > 0
            and (data[cut] & 0xC0) == 0x80
        ):
            cut -= 1

        output.append(
            data[:cut].decode("utf-8")
        )

        data = data[cut:]
        data = b" " + data

    output.append(
        data.decode("utf-8")
    )

    return "\r\n".join(output)


VTZ = """
BEGIN:VTIMEZONE
TZID:Europe/Paris
BEGIN:DAYLIGHT
TZOFFSETFROM:+0100
TZOFFSETTO:+0200
TZNAME:CEST
DTSTART:19700329T020000
RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=-1SU
END:DAYLIGHT
BEGIN:STANDARD
TZOFFSETFROM:+0200
TZOFFSETTO:+0100
TZNAME:CET
DTSTART:19701025T030000
RRULE:FREQ=YEARLY;BYMONTH=10;BYDAY=-1SU
END:STANDARD
END:VTIMEZONE
""".strip().replace("\n", "\r\n")


# =====================================================================
# CALENDAR SCOPE
# =====================================================================

def parse_calendar_scope(scope):
    """
    Convert:

        G1
        G1-A
        G1-B

    into:

        ("G1", None)
        ("G1", "A")
        ("G1", "B")
    """

    match = re.fullmatch(
        r"(G[123])(?:-([AB]))?",
        scope,
        re.IGNORECASE,
    )

    if not match:
        raise ValueError(
            f"Invalid group: {scope}"
        )

    group = match.group(1).upper()
    subgroup = (
        match.group(2).upper()
        if match.group(2)
        else None
    )

    return group, subgroup


def calendar_name(group, subgroup):
    """
    Human-readable calendar name.
    """

    if subgroup:
        return f"ITECH 1 - {group} - {subgroup}"

    return f"ITECH 1 - {group}"


# =====================================================================
# ICS GENERATION
# =====================================================================

def build_ics(
    events,
    group,
    subgroup,
    rooms_by_date,
):
    """
    Build one ICS calendar.

    subgroup:
        None -> complete group
        "A"  -> group A
        "B"  -> group B
    """

    # Fixed timestamp so regenerated calendars remain stable.
    stamp = "20260901T000000Z"

    name = calendar_name(
        group,
        subgroup
    )

    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//planning-itech//FR",
        "CALSCALE:GREGORIAN",
        f"X-WR-CALNAME:{name}",
        f"X-WR-TIMEZONE:{TZ}",
        "REFRESH-INTERVAL;VALUE=DURATION:PT1H",
        "X-PUBLISHED-TTL:PT1H",
        VTZ,
    ]

    count_events = 0
    matched_rooms = 0
    courses_needing_rooms = 0

    for event in events:

        # -----------------------------------------------------------
        # Group filter
        # -----------------------------------------------------------

        if group not in event["groups"]:
            continue

        # -----------------------------------------------------------
        # Subgroup filter
        # -----------------------------------------------------------

        event_subgroups = event.get(
            "subgroups",
            {}
        ).get(
            group,
            set()
        )

        if subgroup is not None:

            if subgroup not in event_subgroups:
                continue

        # -----------------------------------------------------------
        # Event text
        # -----------------------------------------------------------

        event_lines = [
            line.strip()
            for line in event["text"].split("\n")
            if line.strip()
        ]

        title = (
            re.sub(
                r"\s+",
                " ",
                event_lines[0]
            )
            if event_lines
            else "Cours"
        )

        description = [
            re.sub(
                r"\s+",
                " ",
                line
            )
            for line in event_lines[1:]
        ]

        # -----------------------------------------------------------
        # Subgroup information
        # -----------------------------------------------------------

        if subgroup is None:

            # Global calendar:
            # tell the user whether the event is A, B or common.

            if event_subgroups == {"A"}:
                description.append(
                    "Groupe A uniquement"
                )

            elif event_subgroups == {"B"}:
                description.append(
                    "Groupe B uniquement"
                )

            elif event_subgroups == {"A", "B"}:

                if event["partial"]:
                    description.append(
                        "Groupes A et B"
                    )

        else:

            # A/B calendar:
            # only add information for genuinely partial events.
            if event["partial"]:
                description.append(
                    f"Groupe {subgroup}"
                )

        # -----------------------------------------------------------
        # Room matching
        # -----------------------------------------------------------

        rooms = find_rooms(
            event,
            group,
            rooms_by_date,
        )

        location = ""

        if not event["allday"]:

            if rooms:

                location = " / ".join(
                    room["room"]
                    for room in rooms
                )

                if len(rooms) > 1:

                    description.append(
                        "Salles : "
                        + " ; ".join(
                            f"{room['room']}"
                            + (
                                f" ({room['teacher']})"
                                if room["teacher"]
                                else ""
                            )
                            for room in rooms
                        )
                    )

            if (
                not norm(title).startswith("e learning")
                and not re.match(
                    r"\d{1,2}h",
                    title
                )
            ):

                courses_needing_rooms += 1

                if rooms:
                    matched_rooms += 1

        # -----------------------------------------------------------
        # UID
        #
        # Include subgroup so that:
        #
        #   G2
        #   G2-A
        #   G2-B
        #
        # remain separate calendar events.
        # -----------------------------------------------------------

        scope = (
            f"{group}-{subgroup}"
            if subgroup
            else group
        )

        key = (
            f"{scope}|"
            f"{event['date']}|"
            f"{event.get('start', 'allday')}|"
            f"{event.get('end', '')}|"
            f"{norm(title)}"
        )

        uid = (
            hashlib.sha1(
                key.encode("utf-8")
            ).hexdigest()[:20]
            + "@planning-itech"
        )

        # -----------------------------------------------------------
        # VEVENT
        # -----------------------------------------------------------

        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:{uid}",
                f"DTSTAMP:{stamp}",
                f"SUMMARY:{esc(title)}",
            ]
        )

        # -----------------------------------------------------------
        # All-day
        # -----------------------------------------------------------

        if event["allday"]:

            lines.extend(
                [
                    (
                        "DTSTART;VALUE=DATE:"
                        f"{event['date']:%Y%m%d}"
                    ),
                    (
                        "DTEND;VALUE=DATE:"
                        f"{event['end_date']:%Y%m%d}"
                    ),
                    "TRANSP:TRANSPARENT",
                ]
            )

        # -----------------------------------------------------------
        # Timed event
        # -----------------------------------------------------------

        else:

            start = event["start"]
            end = event["end"]

            lines.extend(
                [
                    (
                        f"DTSTART;TZID={TZ}:"
                        f"{event['date']:%Y%m%d}"
                        f"T{start // 60:02d}"
                        f"{start % 60:02d}00"
                    ),
                    (
                        f"DTEND;TZID={TZ}:"
                        f"{event['date']:%Y%m%d}"
                        f"T{end // 60:02d}"
                        f"{end % 60:02d}00"
                    ),
                ]
            )

        # -----------------------------------------------------------
        # Location
        # -----------------------------------------------------------

        if location:
            lines.append(
                f"LOCATION:{esc(location)}"
            )

        # -----------------------------------------------------------
        # Description
        # -----------------------------------------------------------

        if description:

            lines.append(
                "DESCRIPTION:"
                + esc(
                    chr(10).join(description)
                )
            )

        lines.append(
            "END:VEVENT"
        )

        count_events += 1

    # ---------------------------------------------------------------
    # Finish calendar
    # ---------------------------------------------------------------

    lines.append(
        "END:VCALENDAR"
    )

    output = "\r\n".join(
        fold(line)
        if not line.startswith(
            "BEGIN:VTIMEZONE"
        )
        else line
        for line in lines
    ) + "\r\n"

    return (
        output,
        count_events,
        matched_rooms,
        courses_needing_rooms,
    )


# =====================================================================
# GOOGLE SHEETS DOWNLOAD
# =====================================================================

def download(sheet_id, destination):
    """
    Download a public Google Sheet as XLSX.
    """

    url = (
        "https://docs.google.com/spreadsheets/d/"
        f"{sheet_id}/export?format=xlsx"
    )

    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0"
        },
    )

    try:

        data = urllib.request.urlopen(
            request,
            timeout=120
        ).read()

    except Exception as error:

        raise SystemExit(
            "Download failed "
            f"({url}): {error}"
        )

    # XLSX files are ZIP containers.
    if not data.startswith(b"PK"):

        raise SystemExit(
            "Download failed "
            f"({url}): the Google Sheet "
            "does not appear to be public."
        )

    with open(
        destination,
        "wb"
    ) as file:

        file.write(data)


# =====================================================================
# FILE NAME
# =====================================================================

def output_filename(group, subgroup):
    """
    Generate the output filename.

    G1      -> itech1_G1.ics
    G1-A    -> itech1_G1_A.ics
    G1-B    -> itech1_G1_B.ics
    """

    if subgroup:
        return (
            f"itech1_{group}_{subgroup}.ics"
        )

    return (
        f"itech1_{group}.ics"
    )


# =====================================================================
# MAIN
# =====================================================================

def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "planning",
        nargs="?",
    )

    parser.add_argument(
        "salles",
        nargs="?",
    )

    parser.add_argument(
        "--download",
        action="store_true",
        help=(
            "Download the public Google Sheets"
        ),
    )

    parser.add_argument(
        "--groupe",
        choices=[
            "G1",
            "G2",
            "G3",
            "G1-A",
            "G1-B",
            "G2-A",
            "G2-B",
            "G3-A",
            "G3-B",
        ],
        help=(
            "Calendar to generate"
        ),
    )

    parser.add_argument(
        "--tous",
        action="store_true",
        help=(
            "Generate global + A + B calendars "
            "for G1, G2 and G3"
        ),
    )

    parser.add_argument(
        "-o",
        "--out",
        help=(
            "Output file when generating one calendar"
        ),
    )

    parser.add_argument(
        "--outdir",
        default=".",
        help=(
            "Output directory"
        ),
    )

    args = parser.parse_args()

    # ---------------------------------------------------------------
    # Download Google Sheets
    # ---------------------------------------------------------------

    if args.download:

        args.planning = "planning.xlsx"
        args.salles = "salles.xlsx"

        print(
            "Downloading planning Google Sheet..."
        )

        download(
            PLANNING_ID,
            args.planning,
        )

        print(
            "Downloading rooms Google Sheet..."
        )

        download(
            SALLES_ID,
            args.salles,
        )

    # ---------------------------------------------------------------
    # Validate input
    # ---------------------------------------------------------------

    if not (
        args.planning
        and args.salles
    ):

        parser.error(
            "Provide PLANNING.xlsx and SALLES.xlsx, "
            "or use --download."
        )

    if args.out and args.tous:

        parser.error(
            "--out cannot be used with --tous."
        )

    if (
        not args.tous
        and not args.groupe
        and args.out
    ):
        args.groupe = "G1"

    # ---------------------------------------------------------------
    # Read planning
    # ---------------------------------------------------------------

    print(
        "Reading planning..."
    )

    events = read_planning(
        args.planning
    )

    if len(events) < MIN_EVENTS:

        raise SystemExit(
            f"Only {len(events)} events were read. "
            "The spreadsheet structure may have changed. "
            "Nothing was overwritten."
        )

    print(
        f"Planning: {len(events)} events detected."
    )

    # ---------------------------------------------------------------
    # Read rooms
    # ---------------------------------------------------------------

    print(
        "Reading rooms..."
    )

    rooms = read_rooms(
        args.salles
    )

    rooms_by_date = defaultdict(list)

    for room in rooms:
        rooms_by_date[
            room["date"]
        ].append(room)

    print(
        f"Rooms: {len(rooms)} entries detected."
    )

    # ---------------------------------------------------------------
    # Create output directory
    # ---------------------------------------------------------------

    os.makedirs(
        args.outdir,
        exist_ok=True
    )

    # ---------------------------------------------------------------
    # Determine calendars to generate
    # ---------------------------------------------------------------

    if args.tous:

        calendars = [
            ("G1", None),
            ("G1", "A"),
            ("G1", "B"),

            ("G2", None),
            ("G2", "A"),
            ("G2", "B"),

            ("G3", None),
            ("G3", "A"),
            ("G3", "B"),
        ]

    else:

        scope = (
            args.groupe
            or "G1"
        )

        group, subgroup = (
            parse_calendar_scope(scope)
        )

        calendars = [
            (group, subgroup)
        ]

    # ---------------------------------------------------------------
    # Generate calendars
    # ---------------------------------------------------------------

    generated = 0

    for group, subgroup in calendars:

        ics, count, matched, need = build_ics(
            events,
            group,
            subgroup,
            rooms_by_date,
        )

        # -----------------------------------------------------------
        # Output path
        # -----------------------------------------------------------

        if (
            args.out
            and not args.tous
        ):

            output_path = args.out

        else:

            output_path = os.path.join(
                args.outdir,
                output_filename(
                    group,
                    subgroup,
                ),
            )

        # -----------------------------------------------------------
        # Write ICS
        # -----------------------------------------------------------

        with open(
            output_path,
            "w",
            encoding="utf-8",
            newline=""
        ) as file:

            file.write(ics)

        calendar_label = calendar_name(
            group,
            subgroup,
        )

        print(
            f"{output_path} : "
            f"{count} events, "
            f"rooms found for "
            f"{matched}/{need} courses"
        )

        generated += 1

    # ---------------------------------------------------------------
    # Keep repository active
    # ---------------------------------------------------------------

    with open(
        os.path.join(
            args.outdir,
            "last_update.txt"
        ),
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            dt.date.today().isoformat()
            + "\n"
        )

    print(
        f"Generated {generated} calendar(s)."
    )


# =====================================================================
# ENTRY POINT
# =====================================================================

if __name__ == "__main__":
    main()
